from pathlib import Path

import openpyxl
import pytest

from core.missing_input_excel import generate_missing_input_excel, import_missing_input_excel
from core.pipeline import UserInput, apply_overrides_and_revalidate, run_pipeline

FIXED_VALUES = {
    "brand": "트레몰로",
    "manufacturer": "세정",
    "productStatus": "새 상품",
    "stockPerOption": 10,
    "outboundShippingTimeDay": 2,
    "adultOnly": "N",
    "taxType": "Y",
    "parallelImported": "N",
    "overseasPurchased": "N",
}


def touch(path, size=10):
    path.write_bytes(b"x" * size)


@pytest.fixture
def images(tmp_path):
    for i in range(1, 10):
        touch(tmp_path / f"TCNKA26F841104_{i}.jpg")
    return tmp_path


def test_full_workflow_1_to_5(images, tmp_path):
    # 1. 이미지+양식 입력, 아직 색상/가격을 안 줌 -> 자동분석만으로는 막힘
    ready, blocked = run_pipeline(images, {}, FIXED_VALUES)
    assert ready == []
    assert len(blocked) == 1
    product = blocked[0]
    assert set(product.missing_fields) >= {"color", "sale_price"}

    # 2. 추가입력필요.xlsx 생성 (자동입력된 것과 사용자입력 필요한 것 구분)
    missing_path = tmp_path / "추가입력필요.xlsx"
    created = generate_missing_input_excel(blocked, missing_path)
    assert created is True

    wb = openpyxl.load_workbook(missing_path)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "품번" in headers and "색상" in headers and "판매가격" in headers

    # 3. 사용자가 빈칸 보완 (실제로는 엑셀 열어서 타이핑하는 행위)
    col = {h: i + 1 for i, h in enumerate(headers)}
    ws.cell(row=2, column=col["색상"]).value = "네이비"
    ws.cell(row=2, column=col["판매가격"]).value = 29900
    if "소재" in col:
        ws.cell(row=2, column=col["소재"]).value = "폴리에스터 100%"
    if "제조국" in col:
        ws.cell(row=2, column=col["제조국"]).value = "베트남"
    wb.save(missing_path)

    # 3-계속. 재import
    overrides = import_missing_input_excel(missing_path)
    assert overrides["TCNKA26F841104"]["color"] == "네이비"
    assert overrides["TCNKA26F841104"]["sale_price"] == 29900

    # 4. merge + 재검증
    all_products = ready + blocked
    ready2, blocked2 = apply_overrides_and_revalidate(all_products, overrides, FIXED_VALUES)

    # 5. 완벽한 상태로 전환됐는지 확인
    assert blocked2 == []
    assert len(ready2) == 1
    final = ready2[0]
    assert final.color == "네이비"
    assert final.sale_price == 29900
    assert len(final.options) == 4  # KA -> 4사이즈
    assert final.missing_fields == []


def test_still_missing_after_partial_fill_stays_blocked(images, tmp_path):
    ready, blocked = run_pipeline(images, {}, FIXED_VALUES)
    missing_path = tmp_path / "추가입력필요.xlsx"
    generate_missing_input_excel(blocked, missing_path)

    wb = openpyxl.load_workbook(missing_path)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {h: i + 1 for i, h in enumerate(headers)}
    ws.cell(row=2, column=col["색상"]).value = "네이비"
    # 판매가격은 일부러 안 채움
    wb.save(missing_path)

    overrides = import_missing_input_excel(missing_path)
    ready2, blocked2 = apply_overrides_and_revalidate(ready + blocked, overrides, FIXED_VALUES)

    assert ready2 == []
    assert len(blocked2) == 1
    assert "sale_price" in blocked2[0].missing_fields
