import openpyxl

from core.missing_input_excel import generate_missing_input_excel, import_missing_input_excel
from core.models import Product


def make_product(code, color=None, sale_price=None, material=None, missing=None):
    p = Product(product_code=code, color=color, sale_price=sale_price, material=material)
    p.missing_fields = missing or []
    return p


def test_no_file_when_nothing_missing(tmp_path):
    products = [make_product("A", color="블랙", sale_price=10000, material="면", missing=[])]
    out = tmp_path / "추가입력필요.xlsx"
    created = generate_missing_input_excel(products, out)
    assert created is False
    assert not out.exists()


def test_generates_only_needed_columns_and_prefills_known_values(tmp_path):
    products = [
        make_product("A", color=None, sale_price=10000, material="면", missing=["color"]),
        make_product("B", color="블랙", sale_price=None, material=None, missing=["sale_price", "material"]),
    ]
    out = tmp_path / "추가입력필요.xlsx"
    created = generate_missing_input_excel(products, out)
    assert created is True

    wb = openpyxl.load_workbook(out)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    # A는 color만 missing, B는 sale_price+material missing -> 합쳐서 3개 컬럼 필요
    assert headers == ["품번", "색상", "판매가격", "소재"]

    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["A"] == ("A", None, 10000, "면")  # 이미 아는 값(판매가/소재)은 채워져 있고 색상만 빈칸
    assert rows["B"] == ("B", "블랙", None, None)


def test_round_trip_import_after_user_fills_blanks(tmp_path):
    products = [make_product("A", color=None, sale_price=10000, missing=["color"])]
    out = tmp_path / "추가입력필요.xlsx"
    generate_missing_input_excel(products, out)

    # 사용자가 빈칸(색상)을 채워서 저장했다고 가정
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    ws.cell(row=2, column=2).value = "네이비"
    wb.save(out)

    overrides = import_missing_input_excel(out)
    assert overrides == {"A": {"color": "네이비"}}


def test_import_skips_blank_cells(tmp_path):
    products = [
        make_product("A", color=None, sale_price=None, missing=["color", "sale_price"]),
    ]
    out = tmp_path / "추가입력필요.xlsx"
    generate_missing_input_excel(products, out)

    wb = openpyxl.load_workbook(out)
    ws = wb.active
    ws.cell(row=2, column=2).value = "블랙"  # 색상만 채우고 판매가격은 비워둠
    wb.save(out)

    overrides = import_missing_input_excel(out)
    assert overrides == {"A": {"color": "블랙"}}
