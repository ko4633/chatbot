from pathlib import Path

import openpyxl
import pytest

from core.excel_writer import write_products
from core.pipeline import UserInput, run_pipeline

FIXTURES = Path(__file__).resolve().parent.parent / "fixtures"
TEMPLATE = FIXTURES / "coupang_template_v4.6.xlsm"

FIXED_VALUES = {
    "brand": "트레몰로",
    "manufacturer": "세정",
    "productStatus": "새 상품",
    "stockPerOption": 10,
    "outboundShippingTimeDay": 2,
    "adultOnly": "N",
    "taxType": "Y",
    "parallelImported": "N",
    "overseasPurchased": "N",
}


def touch(path, size=10):
    path.write_bytes(b"x" * size)


@pytest.fixture
def sample_images(tmp_path):
    # 정상 상품(사이즈 4개짜리, KA)
    for i in range(1, 10):
        touch(tmp_path / f"TCNKA26F841104_{i}.jpg")
    # DL 상품 (사이즈 8개짜리)
    for i in range(1, 10):
        touch(tmp_path / f"ABCDL26F999999_{i}.jpg")
    # 색상 미입력이라 에러로 막히는 상품
    for i in range(1, 10):
        touch(tmp_path / f"TRNFU5F641609_{i}.jpg")
    return tmp_path


def test_pipeline_end_to_end_writes_real_template(sample_images, tmp_path):
    user_inputs = {
        "TCNKA26F841104": UserInput(color="네이비", sale_price=29900, reference_price=39900),
        "ABCDL26F999999": UserInput(color="차콜", sale_price=59900),
        # TRNFU5F641609 은 색상 없음 -> blocked 로 가야 정상
    }

    ready, blocked = run_pipeline(sample_images, user_inputs, FIXED_VALUES)

    assert {p.product_code for p in ready} == {"TCNKA26F841104", "ABCDL26F999999"}
    assert {p.product_code for p in blocked} == {"TRNFU5F641609"}

    ka_product = next(p for p in ready if p.product_code == "TCNKA26F841104")
    assert len(ka_product.options) == 4  # 095/100/105/110

    dl_product = next(p for p in ready if p.product_code == "ABCDL26F999999")
    assert len(dl_product.options) == 8  # 8사이즈 허리단위

    output_path = tmp_path / "output.xlsm"
    write_products(
        template_path=TEMPLATE,
        output_path=output_path,
        sheet_name="1. 패션잡화",
        category_text="[69657] 패션의류잡화>남성패션>남성잡화>우산>남성2단우산",
        products=ready,
        fixed_values=FIXED_VALUES,
    )

    assert output_path.exists()

    # === 원본 보존 검증 ===
    original_wb = openpyxl.load_workbook(TEMPLATE, keep_vba=True)
    output_wb = openpyxl.load_workbook(output_path, keep_vba=True)

    assert output_wb.sheetnames == original_wb.sheetnames

    # 안 건드린 시트들은 완전히 그대로여야 함
    for sheet_name in ["기본", "2. 식품", "3. 가전", "hidden", "env"]:
        orig_ws = original_wb[sheet_name]
        out_ws = output_wb[sheet_name]
        assert out_ws.max_row == orig_ws.max_row
        assert out_ws.max_column == orig_ws.max_column
        def normalize(v):
            # openpyxl은 저장 시 빈 문자열 셀을 None으로 정규화한다 (우리 코드와 무관한 라이브러리 동작).
            return None if v == "" else v

        for r in range(1, orig_ws.max_row + 1):
            for c in range(1, orig_ws.max_column + 1):
                assert normalize(out_ws.cell(row=r, column=c).value) == normalize(
                    orig_ws.cell(row=r, column=c).value
                ), f"{sheet_name} 시트 {r},{c} 값이 원본과 다름 (건드리면 안 되는 시트)"

    # 대상 시트의 1~4행(안내/헤더 영역)도 그대로여야 함
    orig_target = original_wb["1. 패션잡화"]
    out_target = output_wb["1. 패션잡화"]
    for r in range(1, 5):
        for c in range(1, orig_target.max_column + 1):
            assert out_target.cell(row=r, column=c).value == orig_target.cell(row=r, column=c).value

    # === 실제로 쓴 데이터 검증 (row 순서는 이미지 폴더 스캔 순서에 따라 달라질 수 있어 값 기준으로 찾음) ===
    written_rows = []
    for r in range(5, out_target.max_row + 1):
        product_name = out_target.cell(row=r, column=2).value
        if product_name:
            written_rows.append(r)

    assert len(written_rows) == 4 + 8  # KA(4옵션) + DL(8옵션)

    def find_row(color, size):
        for r in written_rows:
            if out_target.cell(row=r, column=11).value == color and out_target.cell(row=r, column=13).value == size:
                return r
        raise AssertionError(f"색상={color}, 사이즈={size} 옵션행을 찾지 못함")

    ka_row = find_row("네이비", "095")
    assert out_target.cell(row=ka_row, column=1).value == "[69657] 패션의류잡화>남성패션>남성잡화>우산>남성2단우산"
    assert out_target.cell(row=ka_row, column=2).value == "TCNKA26F841104"  # product_name fallback
    assert out_target.cell(row=ka_row, column=7).value == "트레몰로"  # 브랜드
    assert out_target.cell(row=ka_row, column=10).value == "색상"  # 옵션유형1
    assert out_target.cell(row=ka_row, column=12).value == "사이즈"  # 옵션유형2
    assert out_target.cell(row=ka_row, column=62).value == 29900  # 판매가격
    assert out_target.cell(row=ka_row, column=65).value == 10  # 재고수량
    # 이미지 컬럼은 일부러 비워둠: WING 업로드 후 파일명이 달라지므로 사용자가 직접 채워야 함
    assert out_target.cell(row=ka_row, column=104).value is None

    dl_row = find_row("차콜", "076cm(30인치)")
    assert out_target.cell(row=dl_row, column=2).value == "ABCDL26F999999"
    assert out_target.cell(row=dl_row, column=62).value == 59900

    # 기존 샘플로 들어있던 예시행("우산 1" 등)은 지워지고 우리 데이터로 덮였는지
    all_col2_values = [out_target.cell(row=r, column=2).value for r in range(5, out_target.max_row + 1)]
    assert "우산 1" not in all_col2_values
