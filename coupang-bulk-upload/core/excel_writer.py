import zipfile
from pathlib import Path

import openpyxl

from core.models import Product


def _has_real_vba_project(path: str | Path) -> bool:
    """확장자가 .xlsm이어도 실제 vbaProject.bin이 없는 파일이 있다 (예: 참고/예시용 사본).
    이런 파일을 keep_vba=True로 저장하면 openpyxl이 존재하지 않는 vbaProject.bin을 가리키는
    깨진 관계(relationship)를 만들어서 엑셀이 열 때 복구 오류를 띄운다. 그래서 실제로
    vbaProject.bin이 있는지 열어보기 전에 확인하고, 없으면 keep_vba를 강제로 꺼야 한다."""
    if not str(path).lower().endswith(".xlsm"):
        return False
    with zipfile.ZipFile(path) as zf:
        return "xl/vbaProject.bin" in zf.namelist()

HEADER_ROW = 2
GROUP_ROW = 1
DATA_START_ROW = 5


def build_header_map(ws) -> dict[tuple[str, str], int]:
    """(그룹명, 필드명) -> 컬럼번호. 그룹은 1행 병합셀 기준으로 컬럼 범위에 forward-fill."""
    group_by_col: dict[int, str] = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_row == GROUP_ROW and rng.max_row == GROUP_ROW:
            value = ws.cell(row=GROUP_ROW, column=rng.min_col).value
            if value:
                for col in range(rng.min_col, rng.max_col + 1):
                    group_by_col[col] = value

    header_map: dict[tuple[str, str], int] = {}
    for col in range(1, ws.max_column + 1):
        field = ws.cell(row=HEADER_ROW, column=col).value
        if not field:
            continue
        group = group_by_col.get(col, "")
        header_map[(group, field)] = col

    return header_map


def _col(header_map: dict[tuple[str, str], int], group: str, field: str) -> int:
    key = (group, field)
    if key not in header_map:
        raise KeyError(f"엑셀 템플릿에서 컬럼을 찾을 수 없음: group={group!r}, field={field!r}")
    return header_map[key]


def clear_data_rows(ws, start_row: int = DATA_START_ROW, end_row: int | None = None) -> None:
    end_row = end_row or ws.max_row
    for row in range(start_row, end_row + 1):
        for col in range(1, ws.max_column + 1):
            ws.cell(row=row, column=col).value = None


def write_products(
    template_path: str | Path,
    output_path: str | Path,
    sheet_name: str,
    category_text: str,
    products: list[Product],
    fixed_values: dict,
) -> None:
    wb = openpyxl.load_workbook(template_path, keep_vba=_has_real_vba_project(template_path))
    ws = wb[sheet_name]

    header_map = build_header_map(ws)
    clear_data_rows(ws)

    row = DATA_START_ROW
    for product in products:
        for option in product.options:
            ws.cell(row=row, column=_col(header_map, "기본정보", "카테고리")).value = category_text
            ws.cell(row=row, column=_col(header_map, "기본정보", "등록상품명")).value = (
                product.product_name or product.product_code
            )
            ws.cell(row=row, column=_col(header_map, "기본정보", "상품상태")).value = fixed_values["productStatus"]
            ws.cell(row=row, column=_col(header_map, "기본정보", "브랜드")).value = fixed_values["brand"]
            ws.cell(row=row, column=_col(header_map, "기본정보", "제조사")).value = fixed_values["manufacturer"]
            if product.keywords:
                ws.cell(row=row, column=_col(header_map, "기본정보", "검색어")).value = "/".join(product.keywords)

            ws.cell(row=row, column=_col(header_map, "구매옵션", "옵션유형1")).value = "색상"
            ws.cell(row=row, column=_col(header_map, "구매옵션", "옵션값1")).value = option.color
            ws.cell(row=row, column=_col(header_map, "구매옵션", "옵션유형2")).value = "사이즈"
            ws.cell(row=row, column=_col(header_map, "구매옵션", "옵션값2")).value = option.size

            if product.sale_price is not None:
                ws.cell(row=row, column=_col(header_map, "구성 정보", "판매가격")).value = product.sale_price
            if product.reference_price is not None:
                ws.cell(row=row, column=_col(header_map, "구성 정보", "할인율기준가")).value = product.reference_price
            ws.cell(row=row, column=_col(header_map, "구성 정보", "재고수량")).value = option.stock
            ws.cell(row=row, column=_col(header_map, "구성 정보", "출고리드타임")).value = fixed_values[
                "outboundShippingTimeDay"
            ]
            ws.cell(row=row, column=_col(header_map, "구성 정보", "성인상품(19)")).value = fixed_values["adultOnly"]
            ws.cell(row=row, column=_col(header_map, "구성 정보", "과세여부")).value = fixed_values["taxType"]
            ws.cell(row=row, column=_col(header_map, "구성 정보", "병행수입여부")).value = fixed_values[
                "parallelImported"
            ]
            ws.cell(row=row, column=_col(header_map, "구성 정보", "해외구매대행")).value = fixed_values[
                "overseasPurchased"
            ]
            ws.cell(row=row, column=_col(header_map, "구성 정보", "업체상품코드")).value = product.product_code

            # 이미지(대표/추가) 컬럼은 절대 자동으로 채우지 않는다. WING '이미지(파일) 업로드'에서
            # 파일명이 중복되면(예: 상품이미지 탭과 상세설명 탭에 같은 파일을 올리는 경우) WING이
            # 서버에서 예측 불가능한 랜덤 접미사를 붙여 파일명을 바꾼다
            # (실제 확인: "TRHKA5F841156_2.jpg" -> "TRHKA5F841156_2_kbqso.jpg").
            # 이 값은 업로드해보기 전엔 알 수 없으므로, 로컬 파일명을 써넣으면 틀릴 수 있다.
            # 사용자가 WING에서 실제 업로드 후 '복사' 버튼으로 나온 값을 직접 붙여넣어야 한다.

            row += 1

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
