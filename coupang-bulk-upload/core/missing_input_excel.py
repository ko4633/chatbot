from pathlib import Path

import openpyxl

from core.models import Product

# (내부 필드명, 엑셀 헤더 라벨, product에서 값 읽는 함수, 재import시 캐스팅 함수)
FIELD_SPECS = [
    ("color", "색상", lambda p: p.color, str),
    ("sale_price", "판매가격", lambda p: p.sale_price, int),
    ("reference_price", "할인율기준가", lambda p: p.reference_price, int),
    ("material", "소재", lambda p: p.material, str),
    ("country_of_origin", "제조국", lambda p: p.country_of_origin, str),
    ("manufacture_date", "제조년월", lambda p: p.manufacture_date, str),
    ("product_name", "상품명", lambda p: p.product_name, str),
]
FIELD_BY_KEY = {key: spec for spec in FIELD_SPECS for key in [spec[0]]}
LABEL_TO_KEY = {label: key for key, label, _, _ in FIELD_SPECS}


def generate_missing_input_excel(products: list[Product], output_path: str | Path) -> bool:
    """부족한 정보가 있는 상품만 모아 추가입력필요.xlsx 생성.
    반환값 False면 부족한 게 하나도 없어서 파일을 만들 필요가 없었다는 뜻."""
    targets = [p for p in products if p.missing_fields]
    if not targets:
        return False

    needed_keys = []
    for p in targets:
        for key in p.missing_fields:
            if key in FIELD_BY_KEY and key not in needed_keys:
                needed_keys.append(key)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "추가입력필요"

    headers = ["품번"] + [FIELD_BY_KEY[k][1] for k in needed_keys]
    ws.append(headers)

    for p in targets:
        row = [p.product_code]
        for key in needed_keys:
            getter = FIELD_BY_KEY[key][2]
            row.append(getter(p))
        ws.append(row)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return True


def import_missing_input_excel(path: str | Path) -> dict[str, dict]:
    """사용자가 채운 추가입력필요.xlsx -> {품번: {필드명: 값}}. 빈칸은 그대로 빠짐(merge 단계에서 스킵됨)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    columns = list(header_row)
    if not columns or columns[0] != "품번":
        raise ValueError("추가입력필요.xlsx 형식이 아님: 첫 컬럼이 '품번'이어야 함")

    col_keys = [None] + [LABEL_TO_KEY.get(label) for label in columns[1:]]

    result: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        product_code = row[0]
        overrides: dict = {}
        for idx in range(1, len(row)):
            key = col_keys[idx] if idx < len(col_keys) else None
            if key is None:
                continue
            value = row[idx]
            if value in (None, ""):
                continue
            cast = FIELD_BY_KEY[key][3]
            overrides[key] = cast(value)
        result[product_code] = overrides

    return result
